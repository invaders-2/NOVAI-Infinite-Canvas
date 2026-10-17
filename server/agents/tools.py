#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""通用文件与生图工具：read_table / write_html_report / write_xlsx / image_generate。

- 只依赖标准库 + openpyxl（纯 Python，无编译依赖）；
- 运行期上下文通过 contextvar 传入（当前 run 的 run_id / 上传目录 / 产物目录），
  由上下文提供者在执行工具前后 set/reset，工具签名保持 args 单参数，与 AGENT_TOOLS 一致；
- 产物登记走 store.add_artifact，URL 指向受控下载路由；
- image_generate 只做参数校验与参考图解析，真正的生图执行器由 main.py 经
  deps.AGENT_IMAGE_GENERATE 注入（复用 NOVAI 既有 generate_ai_image 链路，不另写 provider 调用）。
"""
from __future__ import annotations

import contextvars
import csv
import datetime
import html as _html
import io
import json
import os
from threading import Lock
from typing import Any, Dict, List, Optional

from server.agents import store
from server.routes.deps import deps as _deps

_CURRENT_RUN = contextvars.ContextVar("novai_agent_run", default=None)

# 同一个 run 里 read_table 的调用上限：防止模型反复读同一张表、把轮次烧光不出产物
_READ_LIMIT = 3
_READ_COUNTS: Dict[str, int] = {}
_READ_LOCK = Lock()


class AgentFileToolError(Exception):
    """文件工具错误（带机器可读 code，交给 agent_execute_step 统一转成统一结构）。"""

    def __init__(self, message: str, code: str = "invalid_args", status_code: int = 400):
        super().__init__(message)
        self.message = str(message)
        self.code = str(code)
        self.status_code = int(status_code)


def set_current_run(ctx: Optional[Dict[str, Any]]):
    return _CURRENT_RUN.set(ctx)


def reset_current_run(token) -> None:
    _CURRENT_RUN.reset(token)


def current_run() -> Dict[str, Any]:
    ctx = _CURRENT_RUN.get()
    if not ctx:
        raise AgentFileToolError("该工具只能在智能体运行中使用", code="no_run_context")
    return ctx


def _uploads_dir() -> str:
    return str(current_run().get("uploads_dir") or "")


def _output_dir() -> str:
    return str(current_run().get("output_dir") or "")


def _run_id() -> str:
    return str(current_run().get("run_id") or "")


def _file_items() -> List[Dict[str, Any]]:
    return current_run().get("files") or []


def _locate(name: str) -> str:
    """按（可省略的）文件名在上传目录里定位文件；找不到就报可读错误。"""
    items = _file_items()
    folder = _uploads_dir()
    wanted = store.safe_filename(name) if str(name or "").strip() else ""
    if not wanted:
        if not items:
            raise AgentFileToolError("本次运行没有上传文件", code="file_not_found")
        wanted = items[0].get("name") or ""
    path = os.path.join(folder, wanted)
    if os.path.isfile(path):
        return path
    # 容错：允许用原始文件名（用户上传时的名字）指代
    for item in items:
        if item.get("original") == str(name or "").strip() and item.get("name"):
            cand = os.path.join(folder, item["name"])
            if os.path.isfile(cand):
                return cand
    names = "、".join([str(i.get("name") or "") for i in items])
    raise AgentFileToolError("找不到文件：%s（本次可用：%s）" % (name or wanted, names or "无"),
                             code="file_not_found", status_code=404)


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
            return str(value)
        return value
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        return str(value)
    if isinstance(value, bytes):
        try:
            return value.decode("utf-8", "replace")
        except Exception:
            return repr(value)
    return str(value)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(_jsonable(value)).strip()


def _to_number(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    if str(value).strip().endswith("%"):
        num = num / 100.0
    return num


def _read_csv_rows(path: str) -> List[List[Any]]:
    raw = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                raw = list(csv.reader(f))
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise AgentFileToolError("无法解码 CSV 文件（已尝试 utf-8 / gb18030）", code="read_failed")
    return raw


def _read_xlsx_sheets(path: str):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise AgentFileToolError("缺少 openpyxl，无法读取 xlsx（请先 pip install openpyxl）",
                                 code="missing_dependency", status_code=500)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise AgentFileToolError("无法打开 xlsx：%s" % exc, code="read_failed")
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([_jsonable(v) for v in row])
            sheets.append({"name": ws.title, "rows": rows})
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return sheets


def _load_sheets(path: str) -> List[Dict[str, Any]]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return [{"name": "CSV", "rows": _read_csv_rows(path)}]
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx_sheets(path)
    if ext == ".xls":
        # 老 xls 需要 xlrd；本期只保证 xlsx/csv
        raise AgentFileToolError("暂不支持 .xls（请另存为 .xlsx 或 .csv）", code="unsupported_format")
    raise AgentFileToolError("不支持的文件类型：%s（只支持 xlsx / csv）" % (ext or "未知"),
                             code="unsupported_format")


def _analyze_rows(rows: List[List[Any]], scan_limit: int = 2000):
    """列名 + 数值列统计 + 低基数列取值分布（供模型据此做真实分析）。"""
    header = [_cell_text(c) for c in (rows[0] if rows else [])]
    width = max([len(r) for r in rows[:scan_limit]] + [len(header), 0])
    while len(header) < width:
        header.append("列%d" % (len(header) + 1))
    columns = []
    used = {}
    for idx, name in enumerate(header):
        label = name or "列%d" % (idx + 1)
        if label in used:
            used[label] += 1
            label = "%s(%d)" % (label, used[label])
        else:
            used[label] = 1
        columns.append(label)
    body = rows[1:]
    numeric_summary = {}
    value_counts = {}
    for idx, label in enumerate(columns):
        values = [_cell_text(r[idx]) if idx < len(r) else "" for r in body[:scan_limit]]
        nums = [_to_number(v) for v in values]
        nums = [n for n in nums if n is not None]
        if nums and len(nums) >= max(2, int(len([v for v in values if v]) * 0.6)):
            numeric_summary[label] = {
                "count": len(nums),
                "sum": round(sum(nums), 6),
                "avg": round(sum(nums) / len(nums), 6),
                "min": round(min(nums), 6),
                "max": round(max(nums), 6),
            }
        distinct = {}
        for v in values:
            if v == "":
                continue
            distinct[v] = distinct.get(v, 0) + 1
        if 0 < len(distinct) <= 12 and len(values) > len(distinct):
            value_counts[label] = [{"value": k, "count": v}
                                   for k, v in sorted(distinct.items(), key=lambda kv: (-kv[1], kv[0]))]
    return columns, numeric_summary, value_counts


def _read_count(run_id: str) -> int:
    with _READ_LOCK:
        return _READ_COUNTS.get(run_id, 0)


def _bump_read(run_id: str) -> int:
    with _READ_LOCK:
        _READ_COUNTS[run_id] = _READ_COUNTS.get(run_id, 0) + 1
        return _READ_COUNTS[run_id]


def read_table(args: Dict[str, Any]) -> Dict[str, Any]:
    args = args if isinstance(args, dict) else {}
    run_id = str(current_run().get("run_id") or "")
    used = _read_count(run_id)
    if used >= _READ_LIMIT:
        raise AgentFileToolError(
            "本次运行已经读取表格 %d 次，数据已经足够。请立即调用 write_html_report / write_xlsx 输出交付文件，不要再重复读表。"
            % used,
            code="read_budget_exhausted")
    path = _locate(args.get("name") or args.get("file") or "")
    sheets = _load_sheets(path)
    if not sheets:
        raise AgentFileToolError("表格里没有可读工作表", code="empty_table")
    wanted = str(args.get("sheet") or "").strip()
    picked = None
    if wanted:
        for sheet in sheets:
            if str(sheet.get("name") or "") == wanted:
                picked = sheet
                break
        if picked is None:
            names = "、".join([str(s.get("name") or "") for s in sheets])
            raise AgentFileToolError("没有工作表：%s（可用：%s）" % (wanted, names), code="sheet_not_found")
    else:
        picked = sheets[0]
    rows = picked.get("rows") or []
    max_rows = args.get("max_rows")
    try:
        max_rows = int(max_rows) if max_rows is not None else 20
    except (TypeError, ValueError):
        max_rows = 20
    max_rows = max(1, min(max_rows, 200))
    try:
        offset = max(0, int(args.get("offset") or 0))
    except (TypeError, ValueError):
        offset = 0
    if not rows:
        _bump_read(run_id)
        return {"file": os.path.basename(path), "sheet": picked.get("name"), "sheets": [s.get("name") for s in sheets],
                "columns": [], "rows": [], "preview": [], "row_count": 0, "returned_rows": 0,
                "truncated": False, "numeric_summary": {}, "value_counts": {}}
    _bump_read(run_id)
    columns, numeric_summary, value_counts = _analyze_rows(rows)
    data_rows = rows[1:]
    total = len(data_rows)
    window = data_rows[offset:offset + max_rows]
    preview = []
    for row in window:
        item = {}
        for idx, label in enumerate(columns):
            item[label] = _cell_text(row[idx]) if idx < len(row) else ""
        preview.append(item)
    return {
        "file": os.path.basename(path),
        "sheet": picked.get("name"),
        "sheets": [s.get("name") for s in sheets],
        "columns": columns,
        "rows": [[_cell_text(c) for c in r] for r in window],
        "preview": preview,
        "row_count": total,
        "returned_rows": len(window),
        "offset": offset,
        "truncated": offset + len(window) < total,
        "numeric_summary": numeric_summary,
        "value_counts": value_counts,
        "hint": "row_count 是数据行总数（不含表头）；preview 就是数据行本身。表格较小时 preview 已包含全部数据，请直接开始分析并输出交付文件，不要重复读表。",
    }


def _render_table(table: Any) -> str:
    if not isinstance(table, dict):
        return ""
    columns = table.get("columns")
    rows = table.get("rows")
    columns = [str(c) for c in columns] if isinstance(columns, list) else []
    rows = rows if isinstance(rows, list) else []
    if not columns and rows and isinstance(rows[0], dict):
        columns = [str(k) for k in rows[0].keys()]
    if not columns:
        return ""
    head = "".join("<th>%s</th>" % _html.escape(c) for c in columns)
    body = []
    for row in rows[:200]:
        if isinstance(row, dict):
            cells = [_html.escape(str(row.get(c, ""))) for c in columns]
        elif isinstance(row, list):
            cells = [_html.escape(str(row[i])) if i < len(row) else "" for i in range(len(columns))]
        else:
            cells = [_html.escape(str(row))]
        body.append("<tr>%s</tr>" % "".join("<td>%s</td>" % c for c in cells))
    return '<div class="tbl-wrap"><table><thead><tr>%s</tr></thead><tbody>%s</tbody></table></div>' % \
        (head, "".join(body))


def _section_title(section: Dict[str, Any]) -> str:
    for key in ("title", "heading", "name"):
        value = section.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def _section_body_text(section: Dict[str, Any]) -> str:
    for key in ("body", "content", "text", "markdown"):
        value = section.get(key)
        if isinstance(value, str) and value.strip():
            return value
    return ""


def _section_has_content(section: Dict[str, Any]) -> bool:
    if str(section.get("html") or "").strip():
        return True
    if _section_body_text(section):
        return True
    items = section.get("items")
    if isinstance(items, list) and any(str(i).strip() for i in items):
        return True
    table = section.get("table")
    if isinstance(table, dict) and (table.get("rows") or table.get("columns")):
        return True
    return False


def _render_section_body(section: Dict[str, Any]) -> str:
    if section.get("html"):
        return str(section["html"])
    parts = []
    body = _section_body_text(section)
    if isinstance(body, str) and body.strip():
        for para in body.split("\n"):
            text = para.strip()
            if not text:
                continue
            if text.startswith(("- ", "* ")):
                parts.append("<p class=\"li\">%s</p>" % _html.escape(text[2:].strip()))
            else:
                parts.append("<p>%s</p>" % _html.escape(text))
    items = section.get("items")
    if isinstance(items, list) and items:
        parts.append("<ul>%s</ul>" % "".join("<li>%s</li>" % _html.escape(str(i)) for i in items if str(i).strip()))
    parts.append(_render_table(section.get("table")))
    return "".join(parts)


def render_html_report(title: str, subtitle: str, sections: List[Dict[str, Any]]) -> str:
    nav, main = [], []
    for idx, section in enumerate(sections or []):
        if not isinstance(section, dict):
            continue
        sid = "sec-%d" % (idx + 1)
        sec_title = _section_title(section) or ("模块 %d" % (idx + 1))
        nav.append('<a href="#%s">%s</a>' % (sid, _html.escape(sec_title)))
        sub = str(section.get("subtitle") or "").strip()
        main.append(
            '<section id="%s" class="card"><h2>%s</h2>%s%s</section>' % (
                sid, _html.escape(sec_title),
                ('<p class="sub">%s</p>' % _html.escape(sub)) if sub else "",
                _render_section_body(section)))
    css = (
        ":root{color-scheme:light}*{box-sizing:border-box}"
        "body{margin:0;font:15px/1.7 -apple-system,BlinkMacSystemFont,'PingFang SC','Microsoft YaHei',sans-serif;"
        "background:#f5f6f8;color:#1f2329}"
        ".layout{display:flex;min-height:100vh}"
        ".side{width:250px;flex:0 0 250px;background:#fff;padding:28px 18px;position:sticky;top:0;height:100vh;overflow:auto}"
        ".side h1{font-size:16px;margin:0 0 6px}.side .sub{color:#8a9099;font-size:12px;margin:0 0 18px}"
        ".side a{display:block;padding:8px 12px;margin:4px 0;border-radius:10px;color:#40454d;text-decoration:none;font-size:14px}"
        ".side a:hover{background:#f0f2f5;color:#111}"
        ".main{flex:1;padding:28px 32px;min-width:0}"
        ".card{background:#fff;border-radius:14px;padding:22px 24px;margin-bottom:18px;box-shadow:0 1px 3px rgba(16,24,40,.06)}"
        ".card h2{font-size:17px;margin:0 0 10px}.card .sub{color:#8a9099;font-size:13px;margin:0 0 12px}"
        "p{margin:6px 0}p.li{position:relative;padding-left:14px}p.li:before{content:'-';position:absolute;left:0;color:#9aa0a6}"
        "ul{margin:8px 0 8px 18px;padding:0}li{margin:4px 0}"
        ".tbl-wrap{overflow:auto;margin-top:10px}"
        "table{border-collapse:collapse;width:100%;font-size:13px}"
        "th,td{padding:8px 10px;text-align:left;vertical-align:top}"
        "th{background:#f0f2f5;font-weight:600;white-space:nowrap}"
        "tbody tr:nth-child(even){background:#fafbfc}"
        ".foot{color:#a4a9b0;font-size:12px;padding:10px 4px 30px}"
    )
    return ("<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
            "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">"
            "<title>%s</title><style>%s</style></head><body><div class=\"layout\">"
            "<nav class=\"side\"><h1>%s</h1><p class=\"sub\">%s</p>%s</nav>"
            "<main class=\"main\">%s<p class=\"foot\">由 NOVAI 智能体生成 · %s</p></main>"
            "</div></body></html>") % (
        _html.escape(title), css, _html.escape(title), _html.escape(subtitle),
        "".join(nav), "".join(main),
        datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))


def write_html_report(args: Dict[str, Any]) -> Dict[str, Any]:
    args = args if isinstance(args, dict) else {}
    title = str(args.get("title") or "智能体分析报告").strip()
    subtitle = str(args.get("subtitle") or "").strip()
    sections = args.get("sections")
    sections = sections if isinstance(sections, list) else []
    raw_html = str(args.get("html") or "").strip()
    if raw_html:
        body_html = raw_html
    else:
        if not sections:
            raise AgentFileToolError("write_html_report 需要 sections 或 html", code="invalid_args")
        body_html = render_html_report(title, subtitle, sections)
    name = store.safe_filename(args.get("name") or "report.html")
    if not name.lower().endswith((".html", ".htm")):
        name = "%s.html" % name
    path = os.path.join(_output_dir(), name)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(body_html)
    size = os.path.getsize(path)
    item = store.add_artifact(_run_id(), "html", name, path)
    return {"kind": "html", "name": name, "url": item.get("url"), "bytes": size,
            "sections": len(sections), "message": "HTML 报告已生成"}


def _normalize_sheets(args: Dict[str, Any]) -> List[Dict[str, Any]]:
    sheets = args.get("sheets")
    if isinstance(sheets, list) and sheets:
        out = [s for s in sheets if isinstance(s, dict)]
        if out:
            return out
    if args.get("rows") is not None or args.get("columns") is not None:
        return [{"name": args.get("sheet") or "Sheet1", "columns": args.get("columns"), "rows": args.get("rows")}]
    return []


def write_xlsx(args: Dict[str, Any]) -> Dict[str, Any]:
    args = args if isinstance(args, dict) else {}
    sheets = _normalize_sheets(args)
    if not sheets:
        raise AgentFileToolError("write_xlsx 需要 sheets（或 columns + rows）", code="invalid_args")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise AgentFileToolError("缺少 openpyxl，无法写出 xlsx（请先 pip install openpyxl）",
                                 code="missing_dependency", status_code=500)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        title = str(sheet.get("name") or "Sheet").strip() or "Sheet"
        ws = wb.create_sheet(title=title[:31])
        columns = sheet.get("columns")
        rows = sheet.get("rows")
        rows = rows if isinstance(rows, list) else []
        columns = [str(c) for c in columns] if isinstance(columns, list) else []
        if not columns and rows and isinstance(rows[0], dict):
            columns = [str(k) for k in rows[0].keys()]
        start = 1
        if columns:
            ws.append(columns)
            for col_idx in range(1, len(columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="F0F2F5")
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            start = 2
        for row in rows:
            if isinstance(row, dict):
                ws.append([_jsonable(row.get(c, "")) for c in columns] if columns else [_jsonable(v) for v in row.values()])
            elif isinstance(row, list):
                ws.append([_jsonable(v) for v in row])
            else:
                ws.append([_jsonable(row)])
        # 列宽按内容粗估，避免打开全是窄列
        max_cols = max([len(columns)] + [len(r) if isinstance(r, list) else 1 for r in rows[:200]] + [1])
        for col_idx in range(1, max_cols + 1):
            width = len(str(columns[col_idx - 1])) if col_idx <= len(columns) else 8
            for row in rows[:200]:
                values = row if isinstance(row, list) else (list(row.values()) if isinstance(row, dict) else [row])
                if col_idx <= len(values):
                    width = max(width, len(str(_jsonable(values[col_idx - 1]))))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(48, max(10, width + 2))
    name = store.safe_filename(args.get("name") or "data.xlsx")
    if not name.lower().endswith(".xlsx"):
        name = "%s.xlsx" % name
    path = os.path.join(_output_dir(), name)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    size = os.path.getsize(path)
    item = store.add_artifact(_run_id(), "xlsx", name, path)
    return {"kind": "xlsx", "name": name, "url": item.get("url"), "bytes": size,
            "sheets": [ws.title for ws in wb.worksheets], "message": "Excel 数据表已生成"}


# --- 生图工具：复用 NOVAI 既有生图链路，执行器由 main.py 经 deps 注入 ---

_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")
_IMAGE_REF_LIMIT = 6
_IMAGE_COUNT_LIMIT = 4


def _is_image_name(name: str) -> bool:
    return os.path.splitext(str(name or ""))[1].lower() in _IMAGE_EXTS


def _resolve_image_ref(ref: Any) -> str:
    """把参考图参数解析成本地绝对路径或 http(s)/data URL；解析不了就明确报错。"""
    if isinstance(ref, dict):
        ref = ref.get("path") or ref.get("url") or ref.get("name") or ""
    text = str(ref or "").strip()
    if not text:
        return ""
    low = text.lower()
    if low.startswith(("http://", "https://", "data:")):
        return text
    if os.path.isabs(text) and os.path.isfile(text):
        return text
    folder = _uploads_dir()
    for cand in (os.path.join(folder, store.safe_filename(text)), os.path.join(folder, text)):
        if os.path.isfile(cand):
            return cand
    for item in _file_items():
        if text in (str(item.get("name") or ""), str(item.get("original") or "")):
            cand = os.path.join(folder, str(item.get("name") or ""))
            if os.path.isfile(cand):
                return cand
    names = "、".join([str(i.get("name") or "") for i in _file_items()])
    raise AgentFileToolError("找不到参考图：%s（本次可用：%s）" % (text, names or "无"),
                             code="file_not_found", status_code=404)


def _default_image_refs() -> List[str]:
    """没显式给参考图时，用本次上传的图片当参考图；没有就纯文生图。"""
    folder = _uploads_dir()
    out = []
    for item in _file_items():
        name = str(item.get("name") or "")
        if not _is_image_name(name):
            continue
        path = os.path.join(folder, name)
        if os.path.isfile(path):
            out.append(path)
    return out


def _image_generator():
    fn = _deps.get("AGENT_IMAGE_GENERATE")
    if fn is None:
        raise AgentFileToolError(
            "本次运行没有接入生图执行器：服务端未注册 AGENT_IMAGE_GENERATE（无法生图，不静默跳过）",
            code="image_backend_missing", status_code=500)
    return fn


async def image_generate(args: Dict[str, Any]) -> Dict[str, Any]:
    """真实调用 NOVAI 既有生图链路生成图片，落到本次运行产物目录并登记为受控产物。"""
    args = args if isinstance(args, dict) else {}
    prompt = str(args.get("prompt") or args.get("text") or "").strip()
    if not prompt:
        raise AgentFileToolError("image_generate 需要 prompt（描述要生成的图）", code="invalid_args")
    refs_arg = args.get("reference_images")
    if refs_arg is None:
        refs_arg = args.get("references")
    if refs_arg is None:
        refs_arg = args.get("images")
    if refs_arg is None:
        refs = _default_image_refs()
    elif isinstance(refs_arg, list):
        refs = []
        for item in refs_arg[:_IMAGE_REF_LIMIT]:
            resolved = _resolve_image_ref(item)
            if resolved:
                refs.append(resolved)
    else:
        raise AgentFileToolError("reference_images 必须是数组", code="invalid_args")
    try:
        count = int(args.get("count") or 1)
    except (TypeError, ValueError):
        count = 1
    count = max(1, min(count, _IMAGE_COUNT_LIMIT))
    result = await _image_generator()({
        "run_id": _run_id(),
        "output_dir": _output_dir(),
        "prompt": prompt,
        "size": str(args.get("size") or "").strip(),
        "quality": str(args.get("quality") or "").strip(),
        "model": str(args.get("model") or "").strip(),
        "provider": str(args.get("provider") or args.get("image_provider") or "").strip(),
        "reference_images": refs,
        "count": count,
        "name_prefix": str(args.get("name") or args.get("name_prefix") or "生成图").strip() or "生成图",
    })
    result = result if isinstance(result, dict) else {}
    images = [i for i in (result.get("images") or []) if isinstance(i, dict)]
    if not images:
        raise AgentFileToolError("生图执行器没有返回任何图片", code="image_empty", status_code=502)
    return {
        "kind": "image",
        "count": len(images),
        "images": images,
        "provider": result.get("provider") or "",
        "model": result.get("model") or "",
        "size": result.get("size") or "",
        "reference_count": len(refs),
        "prompt": prompt[:160] + ("…" if len(prompt) > 160 else ""),
        "message": "已生成 %d 张图片并登记为本次运行产物" % len(images),
    }


# --- 参数校验（与 AGENT_TOOLS 的 validate 约定一致） ---

def _validate_read_table(args):
    args = args if isinstance(args, dict) else {}
    out = {}
    for key in ("name", "file", "sheet"):
        if str(args.get(key) or "").strip():
            out[key] = str(args[key]).strip()
    for key in ("max_rows", "offset"):
        if args.get(key) is not None:
            try:
                out[key] = int(args[key])
            except (TypeError, ValueError):
                raise AgentFileToolError("%s 必须是整数" % key, code="invalid_args")
    return out


def _validate_write_html_report(args):
    args = args if isinstance(args, dict) else {}
    if str(args.get("html") or "").strip():
        return args
    sections = args.get("sections")
    if not isinstance(sections, list) or not sections:
        raise AgentFileToolError("需要 sections（结构化模块数组）或 html（完整 HTML）", code="invalid_args")
    bad = []
    for i, section in enumerate(sections):
        if not isinstance(section, dict):
            bad.append("第 %d 项不是对象" % (i + 1))
            continue
        title = _section_title(section)
        if not title:
            bad.append("第 %d 项缺 title" % (i + 1))
        elif not _section_has_content(section):
            bad.append("第 %d 项（%s）缺正文 body/items/table" % (i + 1, title))
    if bad:
        raise AgentFileToolError(
            "sections 不完整：%s。每个模块必须是 {title, body 或 items 或 table}："
            "title 是模块名，body 写分析结论（可多行），table 写数据表 {columns, rows}。" % "；".join(bad[:5]),
            code="invalid_args")
    return args


def _validate_write_xlsx(args):
    args = args if isinstance(args, dict) else {}
    if not _normalize_sheets(args):
        raise AgentFileToolError("需要 sheets（或 columns + rows）", code="invalid_args")
    return args


def _validate_image_generate(args):
    args = args if isinstance(args, dict) else {}
    if not str(args.get("prompt") or args.get("text") or "").strip():
        raise AgentFileToolError("需要 prompt（描述要生成的图）", code="invalid_args")
    for key in ("reference_images", "references", "images"):
        if args.get(key) is not None and not isinstance(args[key], list):
            raise AgentFileToolError("%s 必须是数组" % key, code="invalid_args")
    return args


def _preview_read_table(args):
    return "读取上传表格：列名 + 前若干行 + 行数"


def _preview_write_html_report(args):
    return "生成 HTML 报告文件并登记为产物"


def _preview_write_xlsx(args):
    return "生成 Excel 数据表并登记为产物"


def _preview_image_generate(args):
    refs = args.get("reference_images") or args.get("references") or args.get("images")
    count = len(refs) if isinstance(refs, list) else "上传图片"
    return "调用 NOVAI 生图链路真实生成图片（参考图 %s，尺寸 %s）" % (count, args.get("size") or "按平台策略")


AGENT_FILE_TOOLS = {
    "read_table": {
        "description": "读取用户上传的 xlsx/csv（只读，不修改文件）：返回真实列名、前 N 行样例、"
                       "数据行总数，以及数值列统计（count/sum/avg/min/max）和低基数列取值分布。"
                       "数据量超过窗口时用 offset 继续读。做任何分析前必须先调用它拿到真实数据。",
        "schema": {"type": "object", "properties": {
            "name": {"type": "string", "description": "文件名（本次只上传一个文件时可省略）"},
            "sheet": {"type": "string", "description": "工作表名（xlsx 多表时可选，默认第一个）"},
            "max_rows": {"type": "integer", "description": "返回的数据行数，默认 20，上限 200"},
            "offset": {"type": "integer", "description": "从第几行开始取（默认 0）"}},
            "required": []},
        "validate": _validate_read_table,
        "run": read_table,
        "preview": _preview_read_table,
        "mutates_canvas": False,
    },
    "image_generate": {
        "description": "真实调用 NOVAI 既有的生图链路生成图片（与画布 generate_image 不同，不建节点），"
                       "产物直接落本次运行目录并登记为受控下载产物。prompt 必填；reference_images 可选"
                       "（本次上传图片的 name / 绝对路径 / http(s) URL；不传时自动用本次上传的图片）；"
                       "size / model / provider 可选，model 缺省用平台 image_models 里的第一个；count 默认 1。",
        "schema": {"type": "object", "properties": {
            "prompt": {"type": "string", "description": "生图提示词（必填）"},
            "reference_images": {"type": "array", "items": {"type": "string"},
                                 "description": "参考图（上传图片名 / 绝对路径 / URL，可选）"},
            "size": {"type": "string", "description": "尺寸，如 1024x1024（可选，按平台策略约束）"},
            "model": {"type": "string", "description": "生图模型（可选，缺省用平台配置的第一个）"},
            "provider": {"type": "string", "description": "生图平台 id（可选）"},
            "quality": {"type": "string", "description": "质量 low/medium/high/auto（可选）"},
            "count": {"type": "integer", "description": "生成张数，默认 1，上限 4"},
            "name": {"type": "string", "description": "产物文件名前缀（可选）"}},
            "required": ["prompt"]},
        "validate": _validate_image_generate,
        "run": image_generate,
        "preview": _preview_image_generate,
        "mutates_canvas": False,
    },
    "write_html_report": {
        "description": "生成 HTML 分析报告文件并登记到本次运行的产物；返回文件名与下载 URL。"
                       "sections 是报告模块数组，每项 {title, subtitle?, body?, items?, table?{columns,rows}}；"
                       "也可以直接给 html（完整 HTML 字符串）。左侧菜单由服务端按 sections 渲染。",
        "schema": {"type": "object", "properties": {
            "title": {"type": "string"},
            "subtitle": {"type": "string"},
            "name": {"type": "string", "description": "输出文件名（可选，默认 report.html）"},
            "sections": {"type": "array", "items": {"type": "object"}, "description": "报告模块数组"},
            "html": {"type": "string", "description": "可选：完整 HTML（给了就不渲染 sections）"}},
            "required": []},
        "validate": _validate_write_html_report,
        "run": write_html_report,
        "preview": _preview_write_html_report,
        "mutates_canvas": False,
    },
    "write_xlsx": {
        "description": "生成 Excel 数据表并登记到本次运行的产物；返回文件名与下载 URL。"
                       "sheets 数组每项 {name, columns?, rows}（rows 为二维数组或对象数组）；"
                       "只有一张表时可直接给 columns + rows。",
        "schema": {"type": "object", "properties": {
            "name": {"type": "string", "description": "输出文件名（可选，默认 data.xlsx）"},
            "sheets": {"type": "array", "items": {"type": "object"}},
            "columns": {"type": "array", "items": {"type": "string"}},
            "rows": {"type": "array", "items": {}}},
            "required": []},
        "validate": _validate_write_xlsx,
        "run": write_xlsx,
        "preview": _preview_write_xlsx,
        "mutates_canvas": False,
    },
}
